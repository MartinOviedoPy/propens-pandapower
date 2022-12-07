# -*- coding: utf-8 -*-
"""
Created on Tue Nov 15 21:40:43 2022

@author: marti
"""

import os
import tempfile

import pandapower.networks as pn
import pandapower as pp
import pandas as pd
#from openpyxl import Workbook
from openpyxl import load_workbook
#from openpyxl.worksheet.table import Table, TableStyleInfo

import Adv_network_only as addnet
import Time_Series_Func as tsf
import Analysis_Func as anal

"""
For testing ADV_Network_Only compatibility, change in line 15~17 and 24
"""
# cases to develop/test the code
#net = pn.create_cigre_network_mv(with_der="all")
#net = pn.case5()
net = addnet.net
#net = pn.panda_four_load_branch()
pp.runpp(net)

network_name = 'Network'
scenario_name = 'Scenario'

# input from fron-end
gen_fuel_tech =[]

output_dir = os.path.join(tempfile.gettempdir(), "time_series_example")
Time_res = tsf.timeseries_example(output_dir)

Sum_Bus_Vol_Under_Data = anal.Anal_Bus_Under(net.res_bus)
Sum_Bus_Vol_Over_Data = anal.Anal_Bus_Over(net.res_bus)
Sum_Trafo_Over_Data = anal.Anal_Trafo_Loading(net.res_trafo)
Sum_Trafo3w_Over_Data = anal.Anal_Trafo3w_Loading(net.res_trafo3w)


# %% read the template and retrieve the sheets
output = 'output_templates/output_template.xlsm'
wb = load_workbook(filename = output, read_only = False, keep_vba = True)

summary_sheet = wb["Summary"]
demand_sheet = wb["Demand"]
generators_sheet = wb["Generation"]
buses_sheet = wb["Buses"]
lines_sheet = wb["Lines"]
trafos_sheet = wb["Trafos"]


# %%
# Preallocate values: number of loads/gen/buses/trafos/lines, and columns in the excel template and their names
loads_number = len(net.load)
generators_number = len(net.gen)
lines_number = len(net.line)
buses_number = len(net.bus)
trafos_number = len(net.trafo)
summary_number = len(Sum_Bus_Vol_Under_Data)

load_column = ['B','C','D','E','F','G','H','I']#,'J']
gen_column = ['B','C','D','E','F','G','H','I','J','K','L','M','N','O','P','Q']
line_column = ['B','C','D','E','F','G','H','I','J','K','L','M','N','O','P','Q','R','S','T','U']
trafo_column = ['B','C','D','E','F','G','H','I','J','K','L','M','N','O','P','Q','R','S','T','U','V']
bus_columm = ['B','C','D','E','F','G','H','I','J','K']
summary_columm = ['B', 'C', 'D' ,'E', 'F']

parameters_net_load = ['zone','bus','vn_kv','in_service']
parameters_res_load = ['p_mw','q_mvar']
parameters_net_gen = ['bus','name','in_service','vm_pu', 'max_p_mw', 'max_q_mvar','min_p_mw','min_q_mvar']
parameters_res_gen = ['p_mw','q_mvar']
parameters_net_line = ['name','from_bus','to_bus', 'in_service','length_km', 'max_i_ka','max_loading_percent', 'parallel', 'std_type']
parameters_res_line = ['p_from_mw', 'q_from_mvar', 'p_to_mw', 'q_to_mvar', 'pl_mw', 'ql_mvar', 'loading_percent']
parameters_net_trafo =['name', 'std_type', 'hv_bus', 'lv_bus', 'vn_hv_kv', 'vn_lv_kv', 'pfe_kw', 'shift_degree','tap_pos', 'parallel', 'in_service' ]
parameters_res_trafo = ['p_from_mw', 'q_from_mvar', 'p_to_mw', 'q_to_mvar', 'pl_mw', 'ql_mvar', 'loading_percent']
parameters_net_bus = ['zone','name','vn_kv','in_service']
parameters_res_bus = ['vm_pu','va_degree','p_mw','q_mvar']

parameters_summary = ['component','Percentage','Extra Info']



steps_number = 2 
initial_line = 4
initial_cell = 'B3'

for step in range(steps_number):
# %% Demand Sheet
    
    if loads_number > 0:
        
        for i in range(len(parameters_net_load)):
            if not parameters_net_load[i] in net.load.keys():  #check the colum of parameter
                net.load[parameters_net_load[i]] = [None]*loads_number #if the columm of parameter is empty, it creat that columm and put value as none
        for i in range(len(parameters_res_load)):
            if not parameters_res_load[i] in net.res_load.keys():
                net.res_load[parameters_res_load[i]] = [None]*loads_number   
        
        
        for i in range(loads_number): # extract the values from power flow results
            load_zone = net.bus.iloc[net.load.loc[i,'bus']]['zone'] #Zone of Bus
            load_index = net.load.index[i]              #Load Index
            net_load = net.load.loc[i,['bus','in_service']].to_list() #Bus Index & In Service
            load_voltage = net.bus.iloc[net.load.loc[i,'bus']]['vn_kv'] #Voltage of Bus that Load is connected
            res_load = net.res_load.loc[i,['p_mw','q_mvar']].to_list()   #P & Q of the load
            load_row = [step] + [load_zone] + [load_index.tolist()] + net_load + [load_voltage] + res_load  #Step/Zone/load Index/Bus Index - In Service/Voltage Level/Active Power-Reactive Power
            
            # write the values in excel table
            for j in range(len(load_row)): # going to all the values in the same line/row
                load_line = i + initial_line + step*loads_number #values start in the row 4 of the sheet
                load_cell = load_column[j] + str(load_line) #update cell reference, from left to right 
                
                if load_row[j] == None:
                    load_value = '---'
                else:
                    load_value = load_row[j]
                
                demand_sheet[load_cell] = load_value # update cell value

    
    # %% Generation's sheet
    
    #check if there is any generators, could be that there is an external grid only
    
    if generators_number > 0:  
    
        for i in range(len(parameters_net_gen)):
            if not parameters_net_gen[i] in net.gen.keys():
                net.gen[parameters_net_gen[i]] = [None]*generators_number
        for i in range(len(parameters_res_gen)):
            if not parameters_res_gen[i] in net.res_gen.keys():
                net.res_gen[parameters_res_gen[i]] = [None]*generators_number        
        if not len(gen_fuel_tech) > 0:
            d = {'fuel': ['---'], 'tech': ['---']}
            gen_fuel_tech = pd.DataFrame(data = d)
            df_aux = pd.DataFrame(data = d)
            for i in range(generators_number):
                gen_fuel_tech = pd.concat([gen_fuel_tech, df_aux],ignore_index=True)         
                    
        for i in range(generators_number): # extract the values from power flow results
            gen_zone = net.bus.iloc[net.gen.loc[i,'bus']]['zone']
            bus_index = net.gen.loc[i,['bus']]
            gen_index = net.gen.index[i]
            
            gen_fuel = gen_fuel_tech.loc[i,'fuel']
            gen_tech = gen_fuel_tech.loc[i,'tech']
            
            gen_voltage = net.bus.iloc[net.gen.loc[i,'bus']]['vn_kv']
            net_gen = net.gen.loc[i,['name','in_service','vm_pu', 'max_p_mw', 'max_q_mvar','min_p_mw','min_q_mvar']].to_list()   
            res_gen = net.res_gen.loc[i,['p_mw','q_mvar']].to_list()
            gen_row = [step] + [gen_zone] + bus_index.tolist() + [gen_index.tolist()] + [gen_fuel] + [gen_tech] + [gen_voltage] + net_gen + res_gen
                
            # write the values in excel table
            for j in range(len(gen_row)): # going to all the values in the same line/row
                gen_line = i + initial_line + step*generators_number #values start in the row 4 of the sheet
                gen_cell = gen_column[j] + str(gen_line) #update cell reference, from left to right 
                
                if gen_row[j] == None:
                    gen_value = '---'
                else:
                    gen_value = gen_row[j]
                    
                generators_sheet[gen_cell] = gen_value # update cell value

    
    # %% Lines' sheet
    
    if lines_number > 0:
       
        for i in range(len(parameters_net_line)):
            if not parameters_net_line[i] in net.line.keys():
                net.line[parameters_net_line[i]] = [None]*lines_number
        for i in range(len(parameters_res_line)):
            if not parameters_res_line[i] in net.res_line.keys():
                net.res_line[parameters_res_line[i]] = [None]*lines_number        
                    
        for i in range(lines_number): # extract the values from power flow results
            line_zone = net.bus.iloc[net.line.loc[i,'from_bus']]['zone']
            line_index = net.line.index[i]
            line_voltage = net.bus.iloc[net.line.loc[i,'from_bus']]['vn_kv']
            net_line = net.line.loc[i,['name','from_bus','to_bus', 'in_service','length_km', 'max_i_ka', 
                                       'max_loading_percent', 'parallel', 'std_type']].to_list()   
            res_line = net.res_line.loc[i,['p_from_mw', 'q_from_mvar', 'p_to_mw', 'q_to_mvar', 'pl_mw',
                                           'ql_mvar', 'loading_percent']].to_list()
            line_row = [step] + [line_zone] + [line_index.tolist()] + [line_voltage] + net_line + res_line
                
            # write the values in excel table
            for j in range(len(line_row)): # going to all the values in the same line/row
                line_line = i + initial_line + step*lines_number#values start in the row 4 of the sheet
                line_cell = line_column[j] + str(line_line) #update cell reference, from left to right 
                
                if line_row[j] == None:
                    line_value = '---'
                else:
                    line_value = line_row[j]
                    
                lines_sheet[line_cell] = line_value # update cell value
           
    
    # %% Trafos' sheet
    
    if trafos_number > 0:
        
        for i in range(len(parameters_net_trafo)):
            if not parameters_net_trafo[i] in net.trafo.keys():
                net.trafo[parameters_net_trafo[i]] = [None]*trafos_number
        for i in range(len(parameters_res_trafo)):
            if not parameters_res_trafo[i] in net.res_trafo.keys():
                net.res_trafo[parameters_res_trafo[i]] = [None]*trafos_number        
                    
        
        for i in range(trafos_number): # extract the values from power flow results
            trafo_zone = net.bus.iloc[net.trafo.loc[i,'hv_bus']]['zone']
            trafo_index = net.trafo.index[i]        
            net_trafo = net.trafo.loc[i,['name', 'std_type', 'hv_bus', 'lv_bus', 'vn_hv_kv', 'vn_lv_kv', 'pfe_kw','shift_degree','tap_pos', 'parallel', 'in_service' ]].to_list()   
            res_trafo = net.res_trafo.loc[i,['p_hv_mw', 'q_hv_mvar', 'p_lv_mw', 'q_lv_mvar','pl_mw', 'ql_mvar', 'loading_percent']].to_list()        
            trafo_row = [step] + [trafo_zone] + [trafo_index.tolist()] + net_trafo + res_trafo
                
            # write the values in excel table
            for j in range(len(trafo_row)): # going to all the values in the same line/row
                trafo_line = i + initial_line + step*trafos_number#values start in the row 4 of the sheet
                trafo_cell = trafo_column[j] + str(trafo_line) #update cell reference, from left to right 
                
                if trafo_row[j] == None:
                    trafo_value = '---'
                else:
                    trafo_value = trafo_row[j]
                    
                trafos_sheet[trafo_cell] = trafo_value # update cell value

        
    # %% Buses Sheet
    
    if buses_number > 0:
        
        for i in range(len(parameters_net_bus)):
            if not parameters_net_bus[i] in net.bus.keys():
                net.bus[parameters_net_bus[i]] = [None]*buses_number
        for i in range(len(parameters_res_bus)):
            if not parameters_res_bus[i] in net.res_bus.keys():
                net.res_bus[parameters_res_bus[i]] = [None]*buses_number        
    
        for i in range(buses_number):
            bus_index = net.bus.index[i]
            net_bus = net.bus.loc[i,['zone','name','vn_kv','in_service']].to_list()
            res_bus = net.res_bus.loc[i,['vm_pu','va_degree','p_mw','q_mvar']].to_list()
            bus_row = [step] + [bus_index.tolist()] + net_bus + res_bus
            
            for j in range(len(bus_row)):
                bus_line = i + initial_line + step*buses_number
                bus_cell = bus_columm[j] + str(bus_line)
                buses_sheet[bus_cell] = bus_row[j] 


# %% Summary Sheet

"""
if there's sth wrong, then Anal function is called.
Anal function is returning certain data base of faults
That data sheet will be written in here.


30Nov Add
I may have to make the Summary_data based on Panda data frame format


30Nov Add_2
When I add all data into one Summary_data, it becacme tuple. 
I need to make different Summary data for each component and add it into one set


6Dec Add

Naming Convention

Sum_Bus_Vol_Under
Sum_Bus_Vol_Over
Sum_Line_Over
Sum_Trafo_Over
Sum_Trafo3w_Over



"""

Sum_Bus_Vol_Under_Data = anal.Anal_Bus_Under(net.res_bus)

if Sum_Bus_Vol_Under_Data is not None :
    for i in range(len(parameters_summary)):
        if not parameters_summary[i] in Sum_Bus_Vol_Under_Data.keys():
            Sum_Bus_Vol_Under_Data[parameters_summary[i]] = [None]*len(Sum_Bus_Vol_Under_Data)
            
    for i in range(len(Sum_Bus_Vol_Under_Data)):
        sum_index = Sum_Bus_Vol_Under_Data.index[i]
        sum_bus_vol_under_index = Sum_Bus_Vol_Under_Data.loc[i,['index']].to_list()
        sum_bus_vol_under_value = Sum_Bus_Vol_Under_Data.loc[i,['value']].to_list()
        sum_bus_vol_under_row = [step] + [sum_index] + sum_bus_vol_under_index + sum_bus_vol_under_value
        
        for j in range(len(sum_bus_vol_under_row)):
            sum_bus_vol_under_line = i + initial_line + 10 + step*len(Sum_Bus_Vol_Under_Data)
            sum_bus_vol_under_cell = summary_columm[j] + str(sum_bus_vol_under_line)
            
            if sum_bus_vol_under_row[j] == None:
                sum_bus_vol_under_value_write = '---'
            else:
                sum_bus_vol_under_value_write = sum_bus_vol_under_row[j]
            
            summary_sheet[sum_bus_vol_under_cell] = sum_bus_vol_under_value_write
            
###############################################################
            
Sum_Bus_Vol_Over_Data = anal.Anal_Bus_Over(net.res_bus)

if Sum_Bus_Vol_Over_Data is not None :
    for i in range(len(parameters_summary)):
        if not parameters_summary[i] in Sum_Bus_Vol_Over_Data.keys():
            Sum_Bus_Vol_Over_Data[parameters_summary[i]] = [None]*len(Sum_Bus_Vol_Over_Data)
            
    for i in range(len(Sum_Bus_Vol_Over_Data)):
        sum_index = Sum_Bus_Vol_Over_Data.index[i]
        sum_bus_vol_over_index = Sum_Bus_Vol_Over_Data.loc[i,['index']].to_list()
        sum_bus_vol_over_value = Sum_Bus_Vol_Over_Data.loc[i,['value']].to_list()
        sum_bus_vol_over_row = [step] + [sum_index] + sum_bus_vol_over_index + sum_bus_vol_over_value
        
        for j in range(len(sum_bus_vol_over_row)):
            sum_bus_vol_over_line = i + initial_line + sum_bus_vol_under_line + step*len(Sum_Bus_Vol_Over_Data)
            sum_bus_vol_over_cell = summary_columm[j] + str(sum_bus_vol_over_line)
            
            if sum_bus_vol_over_row[j] == None:
                sum_bus_vol_over_value_write = '---'
            else:
                sum_bus_vol_over_value_write = sum_bus_vol_over_row[j]
            
            summary_sheet[sum_bus_vol_over_cell] = sum_bus_vol_over_value_write


###############################################################

Sum_Line_Over_Data = anal.Anal_Line_Loading_Better(net.res_line)

if Sum_Line_Over_Data is not None :
    for i in range(len(parameters_summary)):
        if not parameters_summary[i] in Sum_Line_Over_Data.keys():
            Sum_Line_Over_Data[parameters_summary[i]] = [None]*len(Sum_Line_Over_Data)
            
    for i in range(len(Sum_Line_Over_Data)):
        sum_index = Sum_Line_Over_Data.index[i]
        sum_line_over_index = Sum_Line_Over_Data.loc[i,['index']].to_list()
        sum_line_over_value = Sum_Line_Over_Data.loc[i,['value']].to_list()
        sum_line_over_row = [step] + [sum_index] + sum_line_over_index + sum_line_over_value
        
        for j in range(len(sum_line_over_row)):
            sum_line_over_line = i + initial_line + sum_bus_vol_over_line + step*len(Sum_Line_Over_Data)
            sum_line_over_cell = summary_columm[j] + str(sum_line_over_line)
            
            if sum_line_over_row[j] == None:
                sum_line_over_value_write = '---'
            else:
                sum_line_over_value_write = sum_line_over_row[j]
            
            summary_sheet[sum_line_over_cell] = sum_line_over_value_write
            

###############################################################
Sum_Trafo_Over_Data = anal.Anal_Trafo_Loading(net.res_trafo)

if Sum_Trafo_Over_Data is not None :
    for i in range(len(parameters_summary)):
        if not parameters_summary[i] in Sum_Trafo_Over_Data.keys():
            Sum_Trafo_Over_Data[parameters_summary[i]] = [None]*len(Sum_Trafo_Over_Data)
            
    for i in range(len(Sum_Trafo_Over_Data)):
        sum_index = Sum_Trafo_Over_Data.index[i]
        sum_trafo_over_index = Sum_Trafo_Over_Data.loc[i,['index']].to_list()
        sum_trafo_over_value = Sum_Trafo_Over_Data.loc[i,['value']].to_list()
        sum_trafo_over_row = [step] + [sum_index] + sum_trafo_over_index + sum_trafo_over_value
        
        for j in range(len(sum_line_over_row)):
            sum_trafo_over_line = i + initial_line + sum_line_over_line + step*len(Sum_Trafo_Over_Data)
            sum_trafo_over_cell = summary_columm[j] + str(sum_trafo_over_line)
            
            if sum_trafo_over_row[j] == None:
                sum_trafo_over_value_write = '---'
            else:
                sum_trafo_over_value_write = sum_trafo_over_row[j]
            
            summary_sheet[sum_trafo_over_cell] = sum_trafo_over_value_write


###############################################################
Sum_Trafo3w_Over_Data = anal.Anal_Trafo3w_Loading(net.res_trafo3w)

if Sum_Trafo3w_Over_Data is not None :
    for i in range(len(parameters_summary)):
        if not parameters_summary[i] in Sum_Trafo3w_Over_Data.keys():
            Sum_Trafo3w_Over_Data[parameters_summary[i]] = [None]*len(Sum_Trafo3w_Over_Data)
            
    for i in range(len(Sum_Trafo3w_Over_Data)):
        sum_index = Sum_Trafo3w_Over_Data.index[i]
        sum_trafo3w_over_index = Sum_Trafo3w_Over_Data.loc[i,['index']].to_list()
        sum_trafo3w_over_value = Sum_Trafo3w_Over_Data.loc[i,['value']].to_list()
        sum_trafo3w_over_row = [step] + [sum_index] + sum_trafo3w_over_index + sum_trafo3w_over_value
        
        for j in range(len(sum_line_over_row)):
            sum_trafo3w_over_line = i + initial_line + sum_trafo_over_line + step*len(Sum_Trafo3w_Over_Data)
            sum_trafo3w_over_cell = summary_columm[j] + str(sum_trafo3w_over_line)
            
            if sum_trafo3w_over_row[j] == None:
                sum_trafo3w_over_value_write = '---'
            else:
                sum_trafo3w_over_value_write = sum_trafo3w_over_row[j]
            
            summary_sheet[sum_trafo3w_over_cell] = sum_trafo3w_over_value_write
            
            
########################################################################
"""
Maybe it is better to make it a class???

class excel_out_try():
    def __init__(self):
        self.Data_Used

"""

# %% Table reference, and here cell is the last cell added i.e. bottom-right corner of each table

demand_table = demand_sheet.tables["demand_table"]
generation_table = generators_sheet.tables["generation_table"]
trafos_table = trafos_sheet.tables["trafos_table"]
lines_table = lines_sheet.tables["lines_table"]
bus_table = buses_sheet.tables["bus_table"]
        
if loads_number > 0:
    demand_table.ref = initial_cell + ':' + load_cell
if generators_number > 0:  
    generation_table.ref = initial_cell + ':' + gen_cell
if trafos_number > 0:
    trafos_table.ref = initial_cell + ':' + trafo_cell
if lines_number > 0:
    lines_table.ref = initial_cell + ':' + line_cell
if buses_number > 0:
    bus_table.ref = initial_cell + ':' + bus_cell 

# %% save with the topology and scenarios names

filename = 'output_templates/results_' + network_name + '_' + scenario_name + '.xlsm'
wb.save(filename)


"""
Summary_data_Bus_Voltage = anal.Anal_Bus_Under(net.res_bus)

if Summary_data_Bus_Voltage is not None :
        
        for i in range(len(parameters_summary)):
            if not parameters_summary[i] in Summary_data_Bus_Voltage.keys():
                Summary_data_Bus_Voltage[parameters_summary[i]] = [None]*len(Summary_data_Bus_Voltage)
        
        for i in range(len(Summary_data_Bus_Voltage)):
            summary_index = Summary_data_Bus_Voltage.index[i]
            summary_bus_index = Summary_data_Bus_Voltage.loc[i,['index']].to_list()
            summary_bus_value = Summary_data_Bus_Voltage.loc[i, ['value']].to_list()
            Summary_bus_row = [step] + [summary_index] + summary_bus_index + summary_bus_value
            
            for j in range(len(Summary_bus_row)):
                bus_sum_line = i + initial_line + 10 + step*len(Summary_data_Bus_Voltage) #due to button in Macro, I need to shift some a bit down
                bus_sum_cell = summary_columm[j] + str(bus_sum_line)
                
                if Summary_bus_row[j] == None:
                    summary_value = '---'
                else:
                    summary_value = Summary_bus_row[j]
                
                summary_sheet[bus_sum_cell] = summary_value # update cell value
            
"""